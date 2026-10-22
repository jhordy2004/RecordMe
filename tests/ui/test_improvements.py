"""Pruebas unitarias e integración de las 6 mejoras solicitadas."""

import tempfile
from pathlib import Path

import cv2
import docx
import numpy as np
from PySide6.QtCore import Qt
from pytestqt.qtbot import QtBot

from recordme.camera.segmenter import CameraSegmenter
from recordme.config.settings import AppSettings, CameraSettings, TeleprompterSettings
from recordme.core.session import SessionManager
from recordme.resources import get_app_icon
from recordme.ui.camera_settings_dialog import CameraSettingsDialog
from recordme.ui.censor_box import CensorBoxWidget
from recordme.ui.teleprompter import TeleprompterWidget
from recordme.ui.teleprompter_editor import SteppedLinesSlider, TeleprompterEditorDialog
from recordme.ui.toolbar import FloatingToolbar


def test_improvement_1_stepped_lines_slider_snapping(qtbot: QtBot) -> None:
    """Mejora 1: Verifica que el slider de líneas visibles admita hasta 5 líneas,

    no baje de 20% (1 línea) y enganche en 20%, 40%, 60%, 80% y 100% al avanzar y retroceder.
    """
    assert SteppedLinesSlider.snap_percentage(0.0) == 20
    assert SteppedLinesSlider.snap_percentage(10.0) == 20
    assert SteppedLinesSlider.snap_percentage(20.0) == 20

    # 21% al 40% -> 40% (2 líneas)
    assert SteppedLinesSlider.snap_percentage(21.0) == 40
    assert SteppedLinesSlider.snap_percentage(35.0) == 40
    assert SteppedLinesSlider.snap_percentage(40.0) == 40

    # 41% al 60% -> 60% (3 líneas)
    assert SteppedLinesSlider.snap_percentage(41.0) == 60
    assert SteppedLinesSlider.snap_percentage(50.0) == 60
    assert SteppedLinesSlider.snap_percentage(60.0) == 60

    # 61% al 80% -> 80% (4 líneas)
    assert SteppedLinesSlider.snap_percentage(61.0) == 80
    assert SteppedLinesSlider.snap_percentage(75.0) == 80
    assert SteppedLinesSlider.snap_percentage(80.0) == 80

    # 81% al 100% -> 100% (5 líneas)
    assert SteppedLinesSlider.snap_percentage(81.0) == 100
    assert SteppedLinesSlider.snap_percentage(95.0) == 100
    assert SteppedLinesSlider.snap_percentage(100.0) == 100

    # Diálogo de edición completo
    settings = TeleprompterSettings(lines_visible=5)
    dialog = TeleprompterEditorDialog("Texto de prueba", settings)
    qtbot.addWidget(dialog)

    assert dialog.lines_spin.maximum() == 5
    assert dialog.lines_spin.minimum() == 1
    assert dialog.lines_spin.value() == 5
    assert dialog.lines_slider.value() == 100

    # Cambiar spinbox a 1 -> slider a 20
    dialog.lines_spin.setValue(1)
    assert dialog.lines_slider.value() == 20

    # Cambiar slider a valor intermedio 45 -> engancha a 60 (3 líneas)
    dialog.lines_slider.setValue(45)
    assert dialog.lines_slider.value() == 60
    assert dialog.lines_spin.value() == 3

    # Retroceder slider a 30 -> engancha a 40 (2 líneas)
    dialog.lines_slider.setValue(30)
    assert dialog.lines_slider.value() == 40
    assert dialog.lines_spin.value() == 2

    # Verificar settings generados
    cur = dialog.get_current_settings()
    assert cur.lines_visible == 2


def test_improvement_2_file_import_txt_and_word(qtbot: QtBot) -> None:
    """Mejora 2: Verifica que el botón se llame 'Abrir archivo...' y que extraiga texto de .txt y Word (.docx)."""
    settings = TeleprompterSettings()
    dialog = TeleprompterEditorDialog("", settings)
    qtbot.addWidget(dialog)

    assert "Abrir archivo" in dialog.btn_open.text()

    with tempfile.TemporaryDirectory() as tmpdir:
        tmppath = Path(tmpdir)

        # 1. Probar archivo .txt
        txt_file = tmppath / "guion_test.txt"
        txt_file.write_text("Línea uno del guion.\nLínea dos del guion.", encoding="utf-8")
        extracted_txt = dialog._extract_text_from_file(txt_file)
        assert extracted_txt is not None
        assert "Línea uno del guion." in extracted_txt

        # 2. Probar archivo Word .docx
        docx_file = tmppath / "guion_test.docx"
        doc = docx.Document()
        doc.add_paragraph("Párrafo de teleprompter desde Word.")
        doc.add_paragraph("Segunda sección con detalles importantes.")
        doc.save(str(docx_file))

        extracted_docx = dialog._extract_text_from_file(docx_file)
        assert extracted_docx is not None
        assert "Párrafo de teleprompter desde Word." in extracted_docx
        assert "Segunda sección con detalles importantes." in extracted_docx


def test_improvement_3_teleprompter_editor_positioning(qtbot: QtBot) -> None:
    """Mejora 3: Verifica que el diálogo del teleprompter se posicione por debajo de la barra principal."""
    from recordme.app import RecordMeApp

    app_settings = AppSettings()
    session_mgr = SessionManager()
    app = RecordMeApp(session_mgr, app_settings)
    qtbot.addWidget(app.toolbar)

    app.toolbar.show()
    app._open_teleprompter_editor()

    assert app.teleprompter_editor is not None
    assert app.teleprompter_editor.isVisible()
    # Debe estar debajo de la barra superior (y mayor a toolbar.y() + toolbar.height())
    assert app.teleprompter_editor.y() >= app.toolbar.y() + app.toolbar.height()


def test_improvement_4_realtime_background_in_dark_and_cleanup(qtbot: QtBot) -> None:
    """Mejora 4: Verifica que en oscuridad/sin rostro el fondo se mantenga activo,

    y que CameraSettingsDialog limpie la cámara local en done().
    """
    segmenter = CameraSegmenter.get_instance()

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        bg_tmp = f.name
    try:
        # Fondo verde 640x480
        green_bg = np.zeros((480, 640, 3), dtype=np.uint8)
        green_bg[:, :] = (0, 255, 0)
        cv2.imwrite(bg_tmp, green_bg)

        # 1. Fotograma completamente oscuro (luces apagadas en la noche)
        dark_frame = np.zeros((480, 640, 3), dtype=np.uint8)
        res = segmenter.apply_background(dark_frame, bg_tmp)

        # Debe mantener el fondo activo (verde) sin parpadear
        assert res is not None
        assert res.shape == (480, 640, 3)
        # El centro debe ser el fondo verde
        assert res[240, 320, 1] == 255

        # 2. Fotograma con persona iluminada
        lit_frame = np.full((480, 640, 3), 180, dtype=np.uint8)
        cv2.circle(lit_frame, (320, 240), 120, (220, 220, 220), -1)
        res_lit = segmenter.apply_background(lit_frame, bg_tmp)
        assert res_lit is not None
        assert res_lit.shape == (480, 640, 3)

        # 3. Limpieza en CameraSettingsDialog
        cam_settings = CameraSettings()
        dlg = CameraSettingsDialog(cam_settings)
        qtbot.addWidget(dlg)
        dlg.show()
        # Simular cierre/aceptación y verificar liberación de cámara
        dlg.done(1)
        assert dlg._local_source is None
    finally:
        Path(bg_tmp).unlink(missing_ok=True)


def test_improvement_5_cursor_pointers_and_move_cursors(qtbot: QtBot) -> None:
    """Mejora 5: Verifica que los botones tengan cursor pointer y los componentes móviles tengan SizeAllCursor."""
    # 1. Botones de FloatingToolbar y cursor del widget
    session_mgr = SessionManager()
    toolbar = FloatingToolbar(session_mgr)
    qtbot.addWidget(toolbar)
    assert toolbar.cursor().shape() == Qt.CursorShape.SizeAllCursor
    assert toolbar.btn_record.cursor().shape() == Qt.CursorShape.PointingHandCursor
    assert toolbar.btn_mic.cursor().shape() == Qt.CursorShape.PointingHandCursor
    assert toolbar.btn_cam.cursor().shape() == Qt.CursorShape.PointingHandCursor
    assert toolbar.btn_teleprompter.cursor().shape() == Qt.CursorShape.PointingHandCursor

    # 2. Teleprompter flotante
    tp_settings = TeleprompterSettings()
    tp_widget = TeleprompterWidget("Hola", tp_settings)
    qtbot.addWidget(tp_widget)
    assert tp_widget.cursor().shape() == Qt.CursorShape.SizeAllCursor
    assert tp_widget.btn_play_pause.cursor().shape() == Qt.CursorShape.PointingHandCursor
    assert tp_widget.btn_rewind.cursor().shape() == Qt.CursorShape.PointingHandCursor
    assert tp_widget.btn_close.cursor().shape() == Qt.CursorShape.PointingHandCursor

    # 3. Caja de censura
    censor_box = CensorBoxWidget(box_id="test", initial_rect=toolbar.rect())
    qtbot.addWidget(censor_box)
    assert censor_box.cursor().shape() == Qt.CursorShape.SizeAllCursor


def test_improvement_6_app_icon_all_resolutions() -> None:
    """Mejora 6: Verifica que los iconos de src/ico (16, 24, 32, 48, 64, 128, 256) se carguen en el QIcon."""
    icon = get_app_icon()
    assert not icon.isNull()

    avail = [f"{s.width()}x{s.height()}" for s in icon.availableSizes()]
    # Verificar que contiene resoluciones clave
    assert "16x16" in avail
    assert "256x256" in avail
